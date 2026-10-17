"""Parse monthly Belt Tension Report workbooks (<root>/<year>/<NN month>/*.xlsx) into a SQLite file."""
import datetime, glob, hashlib, json, os, re, sqlite3, sys
import openpyxl

MONTHS = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
LEVELS = [1.10, 1.05, 1.03, 1.00, 0.90, 0.80, 0.70]

def num(v):
    return float(v) if isinstance(v, (int, float)) else None

def text(v):
    return v.strip() if isinstance(v, str) and v.strip() not in ("", "-") else None

def parse(path, month_idx, root):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = [list(r) + [None] * 25 for r in ws.iter_rows(values_only=True)]
    date = None
    for r in rows[:4]:
        m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", str(r[0] or ""))
        if m:
            date = f"{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}"
    lines, machines, section, machine = [], [], None, None
    for r in rows:
        c0 = r[0]
        if c0 == "Location (Line no.)": section = "line"; continue
        if c0 == "Location (Machine)": section = "machine"; continue
        if section == "line":
            if not isinstance(c0, str) or not re.match(r"^[A-Z]\d", c0):
                if c0 is None and r[1] is None and r[12] is None and lines: section = None
                continue
            std = num(r[8])
            # CCR placeholders (0.2 t) mean the load cell could not be read
            ccr = lambda v: None if v is None or v < 0.5 else v
            lines.append(dict(
                id=c0, length=num(r[1]), lift_h=num(r[2]), lift_t=num(r[3]), lift_diff=num(r[4]),
                std=std, std_levels=[num(x) for x in r[5:12]],
                before=num(r[12]), after=num(r[13]),
                noload_stop=ccr(num(r[14])), noload_run=ccr(num(r[15])), load=ccr(num(r[16])),
                helix_kn=num(r[19]), helix_ton=num(r[20]), remark=text(r[21]),
                standby=(text(r[21]) or "").lower() == "standby" or r[1] == "-",
            ))
        elif section == "machine":
            if c0 is None and r[2] is None:
                if machines: section = None
                continue
            if isinstance(c0, str) and c0.startswith(("Reported", "Checked")): section = None; continue
            if c0: machine = c0.strip()
            if r[2] is None: continue
            machines.append(dict(
                machine=machine, location=r[2].strip(), std=num(r[8]),
                std_levels=[num(x) for x in r[5:12]],
                left=num(r[12]), right=num(r[13]), remark=text(r[19]),
            ))
    return dict(month=MONTHS[month_idx], month_idx=month_idx, date=date,
                file=os.path.relpath(path, root).replace("\\", "/"),
                lines=lines, machines=machines)

LINE_COLS = ["id", "length", "lift_h", "lift_t", "lift_diff", "std", "before", "after",
             "noload_stop", "noload_run", "load", "helix_kn", "helix_ton", "remark", "standby"]
MACHINE_COLS = ["machine", "location", "std", "left", "right", "remark"]

SCHEMA = """
CREATE TABLE meta (key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE reports (year INTEGER, month_idx INTEGER, month TEXT, date TEXT, file TEXT,
                      PRIMARY KEY (year, month_idx));
CREATE TABLE lines (year INTEGER, month_idx INTEGER, seq INTEGER, id TEXT, length REAL, lift_h REAL,
                    lift_t REAL, lift_diff REAL, std REAL, before REAL, after REAL, noload_stop REAL,
                    noload_run REAL, load REAL, helix_kn REAL, helix_ton REAL, remark TEXT, standby INTEGER,
                    PRIMARY KEY (year, month_idx, seq));
CREATE TABLE machines (year INTEGER, month_idx INTEGER, seq INTEGER, machine TEXT, location TEXT,
                       std REAL, left REAL, right REAL, remark TEXT,
                       PRIMARY KEY (year, month_idx, seq));
"""

def collect(root):
    reports = []
    for p in glob.glob(os.path.join(root, "[0-9][0-9][0-9][0-9]", "*", "*.xlsx")):
        folder = os.path.basename(os.path.dirname(p))
        year = int(os.path.basename(os.path.dirname(os.path.dirname(p))))
        m = re.match(r"(\d{2})", folder)
        if m and not os.path.basename(p).startswith("~$"):
            r = parse(p, int(m[1]) - 1, root)
            r["year"] = year
            reports.append(r)
    reports.sort(key=lambda r: (r["year"], r["month_idx"]))
    return reports

def write_db(reports, out):
    """Rebuild the SQLite file from scratch so it always mirrors the Excel folders."""
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    tmp = out + ".tmp"
    if os.path.exists(tmp):
        os.remove(tmp)
    con = sqlite3.connect(tmp)
    con.executescript(SCHEMA)
    for r in reports:
        key = (r["year"], r["month_idx"])
        con.execute("INSERT INTO reports VALUES (?,?,?,?,?)", key + (r["month"], r["date"], r["file"]))
        for i, l in enumerate(r["lines"]):
            con.execute(f"INSERT INTO lines VALUES ({','.join('?' * (len(LINE_COLS) + 3))})",
                        key + (i,) + tuple(int(l[c]) if c == "standby" else l[c] for c in LINE_COLS))
        for i, mc in enumerate(r["machines"]):
            con.execute(f"INSERT INTO machines VALUES ({','.join('?' * (len(MACHINE_COLS) + 3))})",
                        key + (i,) + tuple(mc[c] for c in MACHINE_COLS))
    # Fingerprint of the content only, so an unchanged rebuild can be detected and not re-pushed
    digest = hashlib.sha256(json.dumps(reports, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    con.executemany("INSERT INTO meta VALUES (?,?)", [
        ("updated_at", datetime.datetime.now().astimezone().isoformat(timespec="seconds")),
        ("content_hash", digest),
    ])
    con.commit()
    con.close()
    os.replace(tmp, out)
    return digest

def read_hash(path):
    if not os.path.exists(path):
        return None
    con = sqlite3.connect(path)
    try:
        row = con.execute("SELECT value FROM meta WHERE key='content_hash'").fetchone()
        return row[0] if row else None
    except sqlite3.Error:
        return None
    finally:
        con.close()

def main():
    root, out = sys.argv[1], sys.argv[2]
    reports = collect(root)
    old = read_hash(out)
    new = write_db(reports, out)
    for r in reports:
        print(r["year"], r["month"], r["date"], len(r["lines"]), "lines", len(r["machines"]), "machine pts")
    print("changed" if old != new else "unchanged")

if __name__ == "__main__":
    main()
